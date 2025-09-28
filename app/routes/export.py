# app/routes/export.py
from typing import Optional, List, Dict
from datetime import datetime as dt, date, timezone, timedelta
import io
from fastapi import APIRouter, Query, HTTPException
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select, and_
from app.core.db import session_scope, init_db
from app.models.orm import EmotionLog

router = APIRouter()
_initialized = False

EMOTION_KEYS = ["楽しい","悲しい","怒り","不安","しんどい","中立"]

def _fmt_local(d: Optional[dt], tzname: str) -> str:
    import zoneinfo
    if d is None:
        return ""
    if d.tzinfo is None:
        d = d.replace(tzinfo=timezone.utc)
    try:
        z = zoneinfo.ZoneInfo(tzname)
    except Exception:
        z = timezone(timedelta(hours=9))
    return d.astimezone(z).strftime("%Y-%m-%d %H:%M")

def _labels_onehot_json(labels: dict) -> str:
    ordered = {k: float(labels.get(k, 0.0)) for k in EMOTION_KEYS}
    return "{" + ", ".join([f"\"{k}\": {v:.1f}" for k, v in ordered.items()]) + "}"

def _parse_iso(value: str, end_of_day: bool = False) -> dt:
    try:
        if "T" in value:
            x = dt.fromisoformat(value.replace("Z","+00:00"))
            if x.tzinfo is not None:
                x = x.astimezone(timezone.utc).replace(tzinfo=None)
            return x
        d = date.fromisoformat(value)
        return dt(d.year, d.month, d.day, 23,59,59) if end_of_day else dt(d.year, d.month, d.day, 0,0,0)
    except Exception:
        raise HTTPException(status_code=422, detail=f"Invalid ISO date/datetime: {value}")

@router.get("", summary="ログをJSON/XLSXでエクスポート")
def export_logs(
    format: str = Query("xlsx", description="xlsx|json"),
    since: Optional[str] = Query(None, description="開始 (YYYY-MM-DD or ISO datetime)"),
    until: Optional[str] = Query(None, description="終了 (YYYY-MM-DD or ISO datetime)"),
    class_id: Optional[str] = Query(default=None, description="クラスID絞り込み"),
    tz: str = Query("Asia/Tokyo", description="表示タイムゾーン（例: Asia/Tokyo）"),
):
    global _initialized
    if not _initialized:
        init_db()
        _initialized = True

    today = date.today()
    if not since:
        since = (today - timedelta(days=7)).isoformat()
    if not until:
        until = today.isoformat()

    start = _parse_iso(since, end_of_day=False)
    end   = _parse_iso(until, end_of_day=True)
    if end < start:
        raise HTTPException(status_code=422, detail="until は since 以上にしてください")

    with session_scope() as s:
        stmt = select(EmotionLog).where(and_(EmotionLog.created_at >= start, EmotionLog.created_at <= end))
        if class_id:
            stmt = stmt.where(EmotionLog.class_id == class_id)
        recs = s.execute(stmt).scalars().all()

    recs_sorted = sorted(recs, key=lambda r: ((r.created_at or dt.min), r.id or 0))

    # JSON
   # app/routes/export.py（JSON 分岐だけ差し替え）
    if format.lower() == "json":
        def iso_z(x: Optional[dt]) -> Optional[str]:
            if x is None: return None
            if x.tzinfo is None:
                x = x.replace(tzinfo=timezone.utc)
            return x.astimezone(timezone.utc).isoformat().replace("+00:00","Z")

        rows_json: List[Dict] = []
        for r in recs_sorted:
            rows_json.append({
                "id": r.id,
                "created_at": iso_z(r.created_at),
                "class_id": r.class_id,
                "emotion": r.emotion,
                "score": float(r.score),
                "relationship_mention": bool(r.relationship_mention),
                "negation_index": float(r.negation_index),
                "avoidance": float(r.avoidance),
                "labels": {k: float((r.labels or {}).get(k, 0.0)) for k in EMOTION_KEYS},
                "topic_tags": list(r.topic_tags or []),
            })
        return JSONResponse(rows_json)  # ★ ← 配列そのまま


    # XLSX
    def autosize(ws):
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = 0
            for cell in col:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(val))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(40, max(10, max_len + 2))

    demo_rows = [{
        "id": r.id,
        "記録日時": _fmt_local(r.created_at, tz),
        "クラス": (r.class_id if r.class_id and str(r.class_id).strip() else "-"),
        "感情": r.emotion,
        "トピック": ";".join(map(str, r.topic_tags or [])),
        "信頼度": round(float(r.score), 3),
    } for r in recs_sorted]

    raw_rows = []
    for r in recs_sorted:
        raw_rows.append({
            "id": r.id,
            "created_at": _fmt_local(r.created_at, tz),
            "class_id": (r.class_id if r.class_id and str(r.class_id).strip() else "-"),
            "emotion_ja": r.emotion,
            "emotion": (str(r.emotion).casefold() if r.emotion else ""),
            "score": round(float(r.score), 3),
            "relationship_mention": bool(r.relationship_mention),
            "negation_index": float(r.negation_index),
            "avoidance": float(r.avoidance),
            "labels_json": _labels_onehot_json(r.labels or {}),
            "topic_tags": ";".join(map(str, r.topic_tags or [])),
        })

    wb = Workbook()

    ws_demo = wb.active
    ws_demo.title = "DEMO"
    demo_headers = ["id","記録日時","クラス","感情","トピック","信頼度"]
    ws_demo.append(demo_headers)
    for row in demo_rows:
        ws_demo.append([row.get(h, "") for h in demo_headers])
    for c in ws_demo[1]:
        c.font = Font(bold=True)
    ws_demo.freeze_panes = "A2"
    ws_demo.auto_filter.ref = ws_demo.dimensions
    autosize(ws_demo)

    ws_raw = wb.create_sheet("RAW")
    raw_headers = ["id","created_at","class_id","emotion_ja","emotion","score",
                   "relationship_mention","negation_index","avoidance",
                   "labels_json","topic_tags"]
    ws_raw.append(raw_headers)
    for row in raw_rows:
        ws_raw.append([row.get(h, "") for h in raw_headers])
    for c in ws_raw[1]:
        c.font = Font(bold=True)
    ws_raw.freeze_panes = "A2"
    ws_raw.auto_filter.ref = ws_raw.dimensions
    autosize(ws_raw)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f'export_demo+raw_{class_id or "all"}_{start.date()}_{end.date()}.xlsx'
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
